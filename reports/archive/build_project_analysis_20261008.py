from pathlib import Path
import sys, json, hashlib, collections, statistics, math
import openpyxl
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'reports'
WORK = OUT / 'archive' / 'project_analysis_20261008'
WORK.mkdir(parents=True, exist_ok=True)

def records(path, sheet):
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rs = list(w[sheet].values)
    w.close()
    return [dict(zip(rs[0], r)) for r in rs[1:] if r[0] is not None]

master = records(ROOT/'Rectifier.xlsx', 'Sheet1')
inventory = records(OUT/'Rectifier_Inventory_188_Sites.xlsx','Sheet1')
settings = records(OUT/'Rectifier_Inventory_188_Sites.xlsx','Settings')
N = len(master)
assert N == 190 and len(inventory) == N and len(settings) == N
assert len({r['ชื่อย่อ'] for r in master}) == N
assert len({r['IP'] for r in master}) == N
assert {r['ชื่อย่อ'] for r in master} == {r['ชื่อย่อ'] for r in inventory}
assert all(r['IP'] for r in master)

def num(x):
    try:
        v=float(x)
        return v if math.isfinite(v) else None
    except (ValueError,TypeError): return None

metrics={}
for key in ['Output Voltage (V)','System Load (A)','Battery Temp (°C)','Ambient Temp (°C)','Estimated Backup Time (h)','Battery Capacity (%)']:
    vs=[num(r[key]) for r in inventory if num(r[key]) is not None]
    metrics[key]={'n':len(vs),'missing':N-len(vs),'min':min(vs),'max':max(vs),'median':statistics.median(vs)}
checks = [('Under Voltage 1 (V)',48),('Under Voltage 2 (V)',46),('LVD 1 Status','Enabled'),('LVD 2 Status','Enabled'),('LVD 1 Disconnect Voltage (V)',43.2),('LVD 2 Disconnect Voltage (V)',43.2),('LVD 1 Reconnect Voltage (V)',49),('LVD 2 Reconnect Voltage (V)',49),('Batt End Test Capacity (%)',70),('Batt End Test Time (min)',30),('Walk-In Status','Enabled'),('Walk-In Time (s)',8)]
audit=[]
for k,target in checks:
    observed=[r for r in settings if r[k] is not None]
    passed=sum((str(r[k])==target if isinstance(target,str) else num(r[k])==target) for r in observed)
    audit.append({'field':k,'target':target,'observed':len(observed),'pass':passed,'mismatch':len(observed)-passed,'missing':N-len(observed)})
w=openpyxl.load_workbook(OUT/'Rectifier_All_Parameters_All_Sites_Latest.xlsx',read_only=True,data_only=True)
raw_count=float_count=bad_float=0
for r in w['All Parameters (Raw)'].iter_rows(min_row=2,values_only=True):
    raw_count+=1
    if r[10]==2:
        float_count+=1
        bad_float+=num(r[12]) is None
w.close()
sources=[]
for path in [ROOT/'Rectifier.xlsx',*OUT.glob('*.xlsx'),ROOT/'README.md',ROOT/'requirements.txt',* (ROOT/'scripts').glob('*.py')]:
    sources.append({'path':str(path.relative_to(ROOT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'modified':path.stat().st_mtime})
evidence={'stations':N,'collection_start':min(r['เวลาเก็บข้อมูล'] for r in inventory),'collection_end':max(r['เวลาเก็บข้อมูล'] for r in inventory),'metrics':metrics,'settings_audit':audit,'raw_count':raw_count,'float_count':float_count,'non_numeric_interpreted_float':bad_float,'sources':sources}
(WORK/'reviewed_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')

FONT='C:/Windows/Fonts/tahoma.ttf'
BOLD='C:/Windows/Fonts/tahomabd.ttf'
def chart_base(title,h=600):
    im=Image.new('RGB',(1500,h),'white'); d=ImageDraw.Draw(im)
    d.text((35,25),title,font=ImageFont.truetype(BOLD,34),fill='#173b56')
    return im,d
im,d=chart_base('ความครบถ้วนของข้อมูลเทียบกับ 190 สถานี',620)
for i,(label,count) in enumerate([('แรงดัน DC',190),('โหลดระบบ',190),('อุณหภูมิห้อง',190),('อุณหภูมิแบตเตอรี่',188),('เวลาสำรองโดยประมาณ',171),('สถานะ LVD 1 / 2',189)]):
    y=110+i*78
    d.text((35,y),label,font=ImageFont.truetype(FONT,29),fill='#263238')
    d.rectangle((440,y+4,1330,y+39),fill='#edf0f2')
    d.rectangle((440,y+4,440+890*count/N,y+39),fill='#297a95')
    d.text((1350,y),str(count),font=ImageFont.truetype(BOLD,29),fill='#173b56')
im.save(WORK/'coverage.png')
im,d=chart_base('อุณหภูมิของสถานีที่ถึงเกณฑ์คัดกรอง 35°C',560)
for i,code in enumerate(['BYD','BU','SB']):
    r=next(r for r in inventory if r['ชื่อย่อ']==code);y=105+i*120
    d.text((45,y+28),code,font=ImageFont.truetype(BOLD,32),fill='#263238')
    for j,(k,c) in enumerate([('Battery Temp (°C)','#297a95'),('Ambient Temp (°C)','#b86e32')]):
        v=num(r[k]);yy=y+j*47
        d.rectangle((210,yy,210+v*24,yy+34),fill=c)
        d.text((225+v*24,yy-2),f'{v:g}°C',font=ImageFont.truetype(FONT,28),fill='#263238')
    d.line((1050,y-8,1050,y+86),fill='#555555',width=2)
d.text((210,490),'แบตเตอรี่ (ฟ้า)   ห้อง (ส้ม)   เส้นตั้ง = 35°C',font=ImageFont.truetype(FONT,28),fill='#263238')
im.save(WORK/'temperature.png')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for name,size in [('Normal',11),('Title',24),('Heading 1',18),('Heading 2',13),('Caption',9)]:
    st=doc.styles[name];st.font.name='Tahoma';st.font.size=Pt(size)
    rp=st.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Tahoma')
    rp.append(rf)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.15
    if name.startswith('Heading'):st.font.color.rgb=RGBColor.from_string('173B56')
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('MEA Rectifier | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    tag=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(tag)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4)
                for run in para.runs:run.font.size=Pt(9)
    return t
def picture(name,caption):
    doc.add_picture(str(WORK/name),width=Inches(6.65));doc.add_paragraph(caption,'Caption')

doc.add_heading('ระบบ Rectifier เก็บข้อมูลครบ 190 สถานี\nแต่ข้อสรุปสุขภาพแบตเตอรี่ยังต้องปรับ',0)
p('รายงานสำรวจโปรเจ็คและวิเคราะห์ข้อมูล | 8 ตุลาคม 2569')
doc.add_heading('บทสรุปสำหรับผู้บริหาร',1)
p('ทะเบียนสถานีและ Inventory ล่าสุดมี 190 สถานี รหัสและ IP ไม่ซ้ำกัน ทุกแถวระบุสถานะเก็บข้อมูล OK โดยเก็บข้อมูลเมื่อ 8 ตุลาคม 2569 เวลา 01:14:43–01:15:12 น. จำนวนจริงต่างจาก README ที่ระบุ 189 สถานี และชื่อไฟล์ Inventory ที่ระบุ 188 Sites [1, 2, 5]')
p('ค่าตั้งสำคัญที่อ่านได้ตรงกับค่ามาตรฐานอ้างอิงในโปรเจ็ค เช่น UV1 = 48 V, UV2 = 46 V, Walk-In = 8 วินาที และ End Test Capacity = 70% ครบ 190 สถานี ส่วน End Test Time ตรง 30 นาทีเพียง 91 สถานี อีก 99 สถานีตั้งไว้ต่างกัน และ TBD ไม่มีข้อมูล LVD ผลผ่านครบชุดตรวจ 12 ฟิลด์จึงมี 90 สถานี [2, 5]')
p('BYD, BU และ SB มีอุณหภูมิห้องหรือแบตเตอรี่ตั้งแต่ 35°C ขึ้นไปตามเกณฑ์คัดกรองของสคริปต์รายงาน ส่วน NBD แสดงกระแสแบตเตอรี่ 2.3 A ขณะสถานะ Float Charge ควรตรวจแนวโน้มและตรวจหน้างานก่อนสรุปความเสื่อม [2, 6]')
p('ข้อจำกัดหลักอยู่ที่ความหมายของข้อมูล: collector เติม Battery Capacity เป็น 100% เมื่อไม่มีค่าที่อ่านได้ และเติม Float Charge กับกระแส 0.0 A เมื่อสถานะแบตเตอรี่หาย นอกจากนี้ตัวแปลงพารามิเตอร์ float นำหน่วยมาใส่ใน Interpreted Value ทำให้ค่า 15,737 จาก 16,682 แถวชนิด float ไม่เป็นตัวเลข [3, 6]')
h('เรื่องที่ควรทำก่อนใช้รายงานตัดสินใจ')
p('ปรับการแยกค่าที่วัดได้กับค่าที่เติมแทน แก้ตัวแปลงพารามิเตอร์ และสร้างรายงานใหม่จากข้อมูลรุ่นเดียวกัน จากนั้นทบทวน End Test Time ของ 99 สถานี ตรวจ TBD และสถานีที่มีสัญญาณเฝ้าระวัง โดยเก็บผลอ่านกลับและหลักฐานการตรวจไว้คู่กับรายงาน')

page('1. โครงสร้างและกระบวนการของโปรเจ็ค')
p('โปรเจ็คใช้ Python เชื่อมต่อ Vertiv / Emerson NetSure NCU ผ่าน HTTP CGI เพื่ออ่าน telemetry, settings และ hardware inventory มีสคริปต์ระดับ scripts/ 16 ไฟล์ และสคริปต์สำรวจย้อนหลังใน scripts/dev/ 28 ไฟล์ ณ วันที่สำรวจ [5, 6]')
table(['ส่วนงาน','ไฟล์หลัก','หน้าที่'],[
('ทะเบียน','Rectifier.xlsx','รายการสถานี รหัส Ring ลำดับ และ IP'),
('ข้อมูลสรุป','collect_rectifier_inventory.py','อ่าน telemetry, settings, inventory แบบหลาย worker'),
('พารามิเตอร์ละเอียด','collect_all_parameters_all_sites.py','สร้าง raw long format และ matrix พร้อมเก็บ archive'),
('รายงาน','rebuild_reports.py','สร้าง abnormal Excel และรายงาน Word'),
('รายงานเพิ่มเติม','generate_comprehensive_word_report.py','สร้างรายงานวิเคราะห์ Word โดยมีข้อความที่กำหนดไว้ในสคริปต์'),
('ตั้งค่า','configure_and_verify_stations.py','login → อ่านก่อน → เขียน CGI → อ่านกลับ'),
('ค่าเสริม','apply_extra_settings.py','เขียนค่าเสริมให้สถานีเป้าหมาย 6 แห่ง')])
h('เส้นทางข้อมูล')
p('Rectifier.xlsx → collector → Inventory / Settings / Raw Parameters → กฎตรวจสอบ → Excel และ Word ใน reports/ ส่วน scripts/dev/, web_assets/ และ config/ เป็นวัสดุอ้างอิงสำหรับพัฒนาการอ่านและตั้งค่า')
h('สิ่งที่ออกแบบไว้แล้ว')
p('collector หลักมี HTTP timeout, retry, ทางเลือก endpoint สำหรับ firmware ต่างรุ่น และใช้ ast.literal_eval แทนการประมวลผลโค้ดจากอุปกรณ์ การตั้งค่าหลักมีการอ่านกลับหลังเขียน และตำแหน่งไฟล์หลายสคริปต์อ้างอิงจาก project root จึงลดปัญหาการรันต่างไดเรกทอรี [6]')
h('ขอบเขตการประเมิน')
p('ผลในรายงานมาจากการตรวจไฟล์และอ่านโค้ด ไม่ใช่การทดสอบการสื่อสารกับอุปกรณ์หรือการทดสอบสมรรถนะแบตเตอรี่ การเก็บข้อมูลครบหมายถึงสถานะของไฟล์ ไม่ใช่การยืนยันความครบถ้วนทุก endpoint')

page('2. ข้อมูลล่าสุดครอบคลุมสถานีใดบ้าง')
picture('coverage.png','ภาพ 1 จำนวนสถานีที่มีค่าใน Inventory / Settings ล่าสุด ช่องที่มีค่าอาจมีค่าที่ collector เติมแทน [2]')
p('ทะเบียนมี 190 รหัสและ 190 IP ที่ไม่ซ้ำกัน รายการรหัสของ Inventory ตรงกับทะเบียนครบทุกสถานี MBD มี IP 10.209.237.20 และมีข้อมูลแล้ว ข้อความเดิมที่ระบุว่า MBD ไม่มี IP จึงไม่ตรงกับชุดล่าสุด [1, 2, 5]')
table(['รายการ','ผลจากไฟล์','ข้อสังเกต'],[
('สถานะเก็บข้อมูล','OK 190 / 190','OK ไม่ครอบคลุมการตรวจทุกช่อง'),
('อุณหภูมิแบตเตอรี่','188 / 190','ขาด RTB และ CPR'),
('เวลาสำรองโดยประมาณ','171 / 190','ขาด 19 สถานี ไม่แทนช่องว่างด้วยศูนย์'),
('สถานะ LVD 1 และ 2','189 / 190','ขาด TBD ทั้งสองรายการ'),
('Config version','1.00: 186; 1.02A: 3; 1.01: 1','ต่างจาก README ที่กล่าวถึง v1.00 เพียง 18 สถานี')])
p('Software version ใน Inventory คือ 5.2.04 จำนวน 184 สถานี, 5.2.05 จำนวน 2 สถานี, 5.2.14B จำนวน 3 สถานี และ 1.3002 จำนวน 1 สถานี การวางแผนรองรับ endpoint ควรใช้การกระจายจริงนี้เป็นฐาน [2]')

page('3. Telemetry และจุดที่ควรตรวจหน้างาน')
table(['ตัวแปร','จำนวนที่มีค่า','ต่ำสุด–สูงสุด','มัธยฐาน'],[
('แรงดัน DC (V)',190,'52.5–54.0','53.3'),('โหลดระบบ (A)',190,'6.2–13.9','6.75'),
('อุณหภูมิแบตเตอรี่ (°C)',188,'18–39','27'),('อุณหภูมิห้อง (°C)',190,'19–40','28'),
('เวลาสำรองโดยประมาณ (h)',171,'10.8–26.4','14.8')])
picture('temperature.png','ภาพ 2 ค่าจาก Inventory วันที่ 8 ตุลาคม 2569 เส้น 35°C เป็นเกณฑ์คัดกรองใน rebuild_reports.py ไม่ใช่ข้อกำหนดรับรองของผู้ผลิต [2, 6]')
p('สถานีอุณหภูมิสูงเป็นกลุ่มเดิม 3 แห่งเมื่อนับแบบไม่ซ้ำกัน: BYD แบตเตอรี่ 39°C / ห้อง 40°C, BU 38°C / 36°C และ SB 35°C / 38°C ควรตรวจระบบปรับอากาศ ตำแหน่ง sensor และอ่านค่าซ้ำเพื่อแยกเหตุชั่วคราวจากปัญหาต่อเนื่อง')
p('NBD มีกระแสแบตเตอรี่ 2.3 A ส่วนอีก 189 สถานีแสดง 0 หรือ 0.0 A ค่าครั้งเดียวและสถานะ Float Charge ยังไม่พอวินิจฉัยแบตเตอรี่เสื่อม เช่นเดียวกับ Estimated Backup Time ซึ่งเป็นค่าประมาณจาก controller และไม่ได้ยืนยันด้วย discharge test [2]')

page('4. ผลตรวจค่าตั้งเทียบฐานอ้างอิง')
p('ใช้ค่าอ้างอิงที่ระบุใน README และสคริปต์ของโปรเจ็ค ไม่ถือเป็นการรับรองข้อกำหนดผู้ผลิตหรือมาตรฐานวิศวกรรมภายนอก เปรียบเทียบค่าตัวเลขหลังแปลงชนิดข้อมูล แยกช่องว่างออกจากผ่านและไม่ผ่าน [2, 5, 6]')
table(['พารามิเตอร์','อ้างอิง','ผ่าน / 190','ขาด'],[
('UV1 / UV2','48 / 46 V','190 / 190',0),
('LVD1 / LVD2 status','Enabled','189 / 190',1),
('LVD1 / LVD2 disconnect','43.2 V','189 / 190',1),
('LVD1 / LVD2 reconnect','49 V','189 / 190',1),
('Walk-In status / time','Enabled / 8 s','190 / 190',0),
('End Test Capacity','70%','190 / 190',0),
('End Test Time','30 min','91 / 190',0)])
p('มี 90 สถานีผ่านครบชุดตรวจ 12 ฟิลด์ (47.4%) อีก 99 สถานีต่างจากอ้างอิงใน End Test Time และ TBD ขาด 6 ฟิลด์ LVD โดย TBD ตั้ง End Test Time 30 นาที จึงจัดเป็นตรวจไม่ครบ กลุ่มทั้งสามไม่ซ้ำกัน')
p('End Test Time ที่ต่างจาก 30 นาทีแบ่งเป็น 15 นาที 83 สถานี, 180 นาที 7 สถานี, 150 นาที 5 สถานี และ 120 นาที 4 สถานี ควรยืนยันเจตนาของค่าตั้งกับแผนทดสอบแบตเตอรี่ก่อนปรับค่าอุปกรณ์')
h('Equalize และการใช้ไฟล์เก่า')
p('Equalize Charge Voltage ใน Inventory ล่าสุดเป็น 53.6 V ครบ 190 สถานี ค่านี้รายงานเป็นข้อเท็จจริงจากไฟล์ ไม่สรุปความเหมาะสมต่อแบตเตอรี่ทุกชนิด สคริปต์ apply_excel_updates.py และ update_reports.py มีการเขียนค่า 53.6 ลง workbook โดยตรง จึงควรเก็บประวัติแยกค่าที่อ่านจากอุปกรณ์กับค่าที่แก้ในรายงาน [2, 6]')
h('ขอบเขตของคำว่า “ผ่าน”')
p('ผลนี้ครอบคลุมเฉพาะพารามิเตอร์ในตาราง ไม่รวมทุกพารามิเตอร์ของ controller ตาราง raw มีข้อมูลหลาย endpoint และบางชื่อซ้ำกัน การเปรียบเทียบเพิ่มเติมควรยึด endpoint, section, equip ID, signal type และ signal ID ร่วมกัน')

page('5. ปัญหาข้อมูลที่ทำให้ข้อสรุปคลาดเคลื่อน')
h('100% ในช่องความจุอาจเป็นค่าที่เติมแทน')
p('collect_rectifier_inventory.py กำหนด batt_cap = "100" เมื่อยังอ่านความจุไม่ได้แต่มีแรงดัน และกำหนด batt_stat = "Float Charge", batt_curr = "0.0" เมื่อสถานะหายแต่มีแรงดัน ชุดล่าสุดแสดงความจุ 100% ทุกสถานี แต่ไฟล์ไม่มี flag บอกว่าค่าใดมาจาก fallback จึงยังนับจำนวนค่าความจุที่วัดจริงไม่ได้ [2, 6]')
h('Interpreted Value ของ float กลายเป็นหน่วย')
p(f'Raw Parameters มี {raw_count:,} แถว ในแถว value_type = 2 จำนวน {float_count:,} แถว มี Interpreted Value ที่แปลงเป็นตัวเลขไม่ได้ {bad_float:,} แถว ({bad_float/float_count*100:.1f}%) ตัวอย่าง KTD / CHARGE.list / Under Voltage 1: Raw Value = 48.0 แต่ Interpreted Value = V [3]')
p('สาเหตุใน extract_parameters_from_dict คือใส่ opts.strip() ลง opt_str สำหรับ float แล้วเลือก opt_str เป็น interpreted_value matrix จึงรับหน่วยไปแทนค่า ควรเก็บ float เป็นตัวเลขและแยกหน่วยใน Options / Unit ซึ่งมีคอลัมน์อยู่แล้ว [6]')
h('เวลาของรายงานไม่ใช่เวลาที่วัด')
p('Inventory มีเวลาเก็บ 8 ตุลาคม แต่ Raw Parameters และรายงานผิดปกติแก้ไขไฟล์เมื่อ 6 ตุลาคม 2569 รายงานผิดปกติยังมีรายการแบตเตอรี่ 2 แถว ขณะที่ Inventory ล่าสุดความจุทุกช่องเป็น 100% การต่างกันนี้ไม่ใช่หลักฐานว่าแบตเตอรี่ฟื้นตัว เพราะความหมายและเวลาของข้อมูลไม่ตรงกัน [2, 3, 4]')
h('การนับสถานีปกติอาจหักซ้ำ')
p('rebuild_reports.py คำนวณสถานีปกติด้วยการลบจำนวนกลุ่มปัญหาหลายกลุ่มจากจำนวนทั้งหมด หากสถานีอยู่ทั้งกลุ่มแบตเตอรี่และอุณหภูมิ จะถูกหักซ้ำ ควรใช้เซตรหัสสถานีของทุกปัญหารวมกันแล้วนับสถานีที่เหลือ [6]')

page('6. โค้ดและกระบวนการที่ควรปรับ')
table(['ประเด็น','หลักฐานในโค้ด','ผลกระทบและแนวแก้'],[
('ค่าที่อ่านไม่ได้ถูกเติม','collector เติม 100%, Float Charge และ 0.0','ใช้ null พร้อม quality_flag และ source_endpoint'),
('นิยาม OK ต่างกัน','collector หลักดู inventory/errors; raw collector ดูว่ามีพารามิเตอร์หรือไม่','แยก connectivity, completeness และ compliance'),
('ตั้งค่าเสริมไม่มี read-back','apply_extra_settings.py เรียก set_parameter แล้วจบ','ใช้กระบวนการอ่านก่อน/หลัง และบันทึก mismatch'),
('ข้ามพารามิเตอร์เงียบ','configure_and_verify_stations.py ใช้ continue เมื่อไม่พบ row','คืนสถานะ MISSING_PARAMETER และห้ามสรุปครบมาตรฐาน'),
('เป้าหมายระบุในโค้ด','TARGET_STATIONS กำหนด code/IP ไว้ 6 แห่ง','เทียบคู่ code/IP กับทะเบียนก่อนส่งทุกคำสั่ง'),
('ข้อมูลลับอยู่ใน source/log','login มี credential literal และพิมพ์ session ID','รับค่าจาก environment หรือ secret store; ปิดการพิมพ์ token'),
('dependency ไม่ครบ','requirements.txt มี openpyxl/playwright; รายงาน import docx','เพิ่ม python-docx และระบุ dependency ของ workflow'),
('กฎและข้อความกำหนดตายตัว','รายงานมีเงื่อนไข NBD, ค่า 94% และข้อความความเสื่อม','ให้ข้อความเกิดจากข้อมูล/เกณฑ์ที่มีเหตุผลและหลักฐาน')])
p('parser หลักจัดการ marker, comment และ pseudo-JSON ผ่าน json.loads/ast.literal_eval ส่วน raw collector มี parser ของตนเอง การรวมเป็นฟังก์ชันเดียวพร้อมตัวอย่าง firmware จะช่วยลดผลอ่านต่างกัน เงื่อนไข if not value ควรแยกค่าศูนย์จริงจากค่าว่างด้วย [6]')
p('การเชื่อมต่อ CGI ใช้ HTTP ในเครือข่ายสถานี จึงควรควบคุมเครื่องที่เข้าถึงได้และปกป้องไฟล์หลักฐานที่มี IP กับ serial number ข้อเสนอด้านบัญชีผู้ใช้และการเปลี่ยนค่าเป็นงานปรับปรุงกระบวนการ ยังไม่ได้ดำเนินการในรายงานนี้')

page('7. ลำดับปรับปรุงและเกณฑ์ตรวจรับ')
table(['ลำดับ','งาน','หลักฐานที่ใช้ปิดงาน'],[
('ก่อนสร้างรายงานรอบถัดไป','แก้ float interpretation และ fallback metadata','fixture ได้ค่า 48.0 แยกหน่วย V; ค่าหายเป็น null'),
('ก่อนสรุปผลทั้งระบบ','ปรับจำนวนสถานีและนิยามผลตรวจ; ทบทวน End Test Time 99 สถานี','รายการผ่าน 90 / ต่างค่า 99 / ตรวจไม่ครบ 1; ยืนยันค่าเป้าหมาย'),
('ตรวจหน้างานตามสัญญาณ','BYD, BU, SB, NBD และ TBD; ตรวจ sensor RTB/CPR','อ่านซ้ำพร้อมเวลา และผลตรวจที่ระบุสถานี'),
('ก่อนตั้งค่าอุปกรณ์รอบถัดไป','รวม updater ให้ตรวจ code/IP และอ่านกลับ','บันทึก before/target/after ต่อ parameter; missing ไม่เป็น success'),
('ทำให้รายงานทำซ้ำได้','สร้างทุก output จาก snapshot เดียว เก็บ hash และเวลากำเนิด','ผลนับตรงกัน; รันซ้ำไม่แก้ raw snapshot; นับปัญหาแบบเซต'),
('ทดสอบซอฟต์แวร์','เพิ่ม fixture parser, enum, float, missing, overlapping issues','ตรวจกรณี NaN/undefined/comment/0 และผลสรุปที่ไม่หักซ้ำ')])
h('แหล่งข้อมูลและวิธีตรวจ')
p('[1] Rectifier.xlsx / Sheet1: 190 สถานี ตรวจจำนวนรหัส, IP, ความซ้ำ และจับคู่รหัสกับ Inventory')
p('[2] reports/Rectifier_Inventory_188_Sites.xlsx / Sheet1, Settings, Inventory: telemetry และค่าตั้ง เวลาเก็บในแถว 2026-10-08T01:14:43+07:00 ถึง 01:15:12+07:00')
p('[3] reports/Rectifier_All_Parameters_All_Sites_Latest.xlsx / Station Summary, Settings Matrix (All Sites), All Parameters (Raw), Hardware Inventory: เก็บแยกจาก Inventory; file modified 6 ตุลาคม ใช้วันแก้ไขเป็นข้อมูลไฟล์ ไม่ใช่ยืนยันวันวัดทุกแถว')
p('[4] reports/Rectifier_Abnormal_Report.xlsx: ตรวจจำนวนรายการเทียบกับชุดล่าสุด ไม่ใช้เป็นสถานะปัจจุบันเมื่อข้อมูลต่างรุ่น')
p('[5] README.md และ AGENTS.md: ขอบเขตสถาปัตยกรรมและค่าอ้างอิงภายในโปรเจ็ค')
p('[6] scripts/collect_rectifier_inventory.py, collect_all_parameters_all_sites.py, rebuild_reports.py, configure_and_verify_stations.py, apply_extra_settings.py, apply_excel_updates.py, update_reports.py และ requirements.txt: อ่านโค้ดโดยไม่รันคำสั่งกับ controller')
p('หลักฐานคำนวณและ SHA-256 ของแหล่งข้อมูลอยู่ที่ reports/archive/project_analysis_20261008/reviewed_evidence.json เกณฑ์คำนวณใช้จำนวนสถานี 190 เป็นฐาน ช่องที่หายไม่เป็นศูนย์ และแยกผ่าน/ผิดค่า/ตรวจไม่ครบ')

doc.core_properties.title='รายงานสำรวจและวิเคราะห์โปรเจ็ค MEA Rectifier'
doc.core_properties.subject='Project review and evidence-based telemetry analysis'
doc.core_properties.author=''
file=OUT/'Rectifier_Project_Analysis_20261008.docx'
doc.save(file)
assert sum(x['mismatch'] for x in audit)==99
assert sum(all((str(r[k])==v if isinstance(v,str) else num(r[k])==v) for k,v in checks) for r in settings)==90
print(json.dumps({'output':str(file),'stations':N,'metrics':metrics,'raw_count':raw_count,'float_non_numeric':bad_float},ensure_ascii=False))
