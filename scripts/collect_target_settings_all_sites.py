# -*- coding: utf-8 -*-
"""
scripts/collect_target_settings_all_sites.py — Live Collector for Temperature, Rectifiers, Battery Charge, and System Settings

Pulls all settings live across all 189 MEA substations for:
  - Battery Charge (CHARGE)
  - Rectifiers (RECT)
  - Temperature (TEMPERATURE)
  - System (PowerSystem / sys_setting)

Generates multi-sheet Excel with matrices and detailed parameter breakdown.
"""
from __future__ import annotations
import sys
import time
import json
import re
import ast
import urllib.request
import urllib.error
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', line_buffering=True)
    except Exception:
        pass

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_MARKER = '"%__data_end__%"'


def parse_device_data(raw: bytes) -> dict[str, Any]:
    text = raw.decode("utf-8-sig", errors="replace")
    text = re.sub(r"/\*[\s\S]*?\*/|//.*", "", text).strip()
    text = text.split(DATA_MARKER, 1)[0].strip()
    try:
        return json.loads(text)
    except Exception:
        python_text = re.sub(r"\btrue\b", "True", text)
        python_text = re.sub(r"\bfalse\b", "False", python_text)
        python_text = re.sub(r"\bnull\b", "None", python_text)
        python_text = re.sub(r"\bundefined\b", "None", python_text)
        python_text = re.sub(r"\bnan\b", "None", python_text, flags=re.IGNORECASE)
        try:
            parsed = ast.literal_eval(python_text)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            pass
    return {}


def fetch_json(ip: str, path: str, timeout: float = 3.5) -> dict[str, Any]:
    url = f"http://{ip}{path}?_={int(time.time() * 1000)}"
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "MEA-Rectifier/1.0", "Cache-Control": "no-cache"}
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return parse_device_data(resp.read())
    except Exception:
        return {}


def interpret_param_row(item: list) -> dict[str, Any]:
    """Interprets a single parameter array [name, equip_id, sig_type, sig_id, level, ts, max, min, vt, val, opts]."""
    raw_name = item[0]
    name = raw_name[0] if isinstance(raw_name, list) else str(raw_name)
    eq_id = item[1] if len(item) > 1 else None
    sig_type = item[2] if len(item) > 2 else None
    sig_id = item[3] if len(item) > 3 else None
    val_type = item[8] if len(item) > 8 else None
    val = item[9] if len(item) > 9 else None
    opts = item[10] if len(item) > 10 else None

    interpreted = str(val) if val is not None else ""
    unit_str = ""

    if val_type == 5 and isinstance(opts, list):
        try:
            idx = int(val)
            if 0 <= idx < len(opts):
                o = opts[idx]
                interpreted = o[0] if isinstance(o, list) else str(o)
        except Exception:
            pass
        unit_str = "enum"
    elif val_type == 2 and isinstance(opts, str):
        unit_str = opts.strip()
    elif isinstance(opts, str):
        unit_str = opts.strip()

    return {
        "name": name.strip(),
        "raw_value": val,
        "interpreted_value": interpreted.strip(),
        "unit": unit_str,
        "val_type": val_type,
        "equip_id": eq_id,
        "sig_type": sig_type,
        "sig_id": sig_id,
    }


def collect_site_settings(site: dict[str, Any], timeout: float = 3.5) -> dict[str, Any]:
    code = site.get("code", "")
    name = site.get("name", "")
    ring = site.get("ring", "")
    no = site.get("no", "")
    ip = str(site.get("ip", "")).strip()

    if not ip or ip.lower() in ("none", "nan", "ไม่มี", "-", ""):
        return {
            "site": site,
            "status": "NO_IP",
            "settings": {},
            "all_params": [],
            "error": "No IP Address"
        }

    d_set = fetch_json(ip, "/var/datas/data.setting.html", timeout=timeout)
    d_sys = fetch_json(ip, "/var/datas/data.sys_setting.html", timeout=timeout)

    if not d_set and not d_sys:
        return {
            "site": site,
            "status": "OFFLINE",
            "settings": {},
            "all_params": [],
            "error": "Timeout / Connection Error"
        }

    categories = {
        "Battery Charge": {},
        "Rectifiers": {},
        "Temperature": {},
        "System": {}
    }
    flat_params = []

    # 1. Parse data.setting.html
    cat_map = {
        "CHARGE": "Battery Charge",
        "RECT": "Rectifiers",
        "TEMPERATURE": "Temperature"
    }

    for src_key, target_cat in cat_map.items():
        sub_list = d_set.get(src_key, {}).get("list", [])
        for item in sub_list:
            if isinstance(item, list) and len(item) >= 10:
                parsed = interpret_param_row(item)
                param_name = parsed["name"]
                categories[target_cat][param_name] = parsed["interpreted_value"]
                flat_params.append({
                    "category": target_cat,
                    **parsed
                })

    # 2. Parse data.sys_setting.html -> System
    ps = d_sys.get("PowerSystem", {})
    def walk_sys(node, group_label=""):
        if isinstance(node, dict):
            for k, v in node.items():
                lbl = k[0] if isinstance(k, list) else str(k)
                walk_sys(v, lbl)
        elif isinstance(node, list):
            for itm in node:
                if isinstance(itm, list) and len(itm) >= 10 and isinstance(itm[1], int) and isinstance(itm[2], int):
                    parsed = interpret_param_row(itm)
                    p_name = parsed["name"]
                    categories["System"][p_name] = parsed["interpreted_value"]
                    flat_params.append({
                        "category": "System",
                        **parsed
                    })
                elif isinstance(itm, (dict, list)):
                    walk_sys(itm, group_label)

    walk_sys(ps)

    return {
        "site": site,
        "status": "ONLINE",
        "settings": categories,
        "all_params": flat_params,
        "error": ""
    }


def main():
    print("=" * 70)
    print("MEA RECTIFIER SETTINGS LIVE COLLECTOR (ALL SITES)")
    print("Categories: Temperature, Rectifiers, Battery Charge, System")
    print(f"Timestamp : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 70)

    # 1. Load Station List
    excel_path = PROJECT_ROOT / "Rectifier.xlsx"
    if not excel_path.exists():
        print(f"[-] Rectifier.xlsx not found at {excel_path}")
        return

    wb_in = openpyxl.load_workbook(excel_path, read_only=True)
    ws_in = wb_in.active

    sites = []
    for row in ws_in.iter_rows(values_only=True):
        if not row or not row[0] or str(row[0]).strip().lower() in ("ชื่อย่อ", "code", "station"):
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        ring = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        no = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        ip = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        sites.append({"code": code, "name": name, "ring": ring, "no": no, "ip": ip})
    wb_in.close()

    print(f"[*] Loaded {len(sites)} stations from Rectifier.xlsx")
    print(f"[*] Dispatching concurrent HTTP requests (workers=25, timeout=3.5s)...")

    t_start = time.time()
    results = []

    with ThreadPoolExecutor(max_workers=25) as executor:
        future_to_site = {executor.submit(collect_site_settings, s): s for s in sites}
        for future in as_completed(future_to_site):
            results.append(future.result())

    # Sort results by original site order
    site_order = {s["code"]: idx for idx, s in enumerate(sites)}
    results.sort(key=lambda r: site_order.get(r["site"]["code"], 999))

    elapsed = time.time() - t_start
    online_count = sum(1 for r in results if r["status"] == "ONLINE")
    offline_count = sum(1 for r in results if r["status"] == "OFFLINE")
    no_ip_count = sum(1 for r in results if r["status"] == "NO_IP")

    print(f"\n[+] Collection completed in {elapsed:.2f} seconds!")
    print(f"    - Online  : {online_count}/{len(sites)}")
    print(f"    - Offline : {offline_count}")
    print(f"    - No IP   : {no_ip_count}\n")

    # Collect unique parameter names per category for matrix building
    cat_keys = ["Battery Charge", "Rectifiers", "Temperature", "System"]
    cat_params = {c: set() for c in cat_keys}
    all_params_order = []

    for r in results:
        if r["status"] == "ONLINE":
            for c in cat_keys:
                cat_params[c].update(r["settings"].get(c, {}).keys())

    sorted_cat_params = {c: sorted(list(cat_params[c])) for c in cat_keys}

    # 2. Build Excel Workbook
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill_brand = PatternFill(start_color="F26522", end_color="F26522", fill_type="solid") # MEA Orange
    header_fill_sub = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Dark Blue
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    base_cols = ["No.", "Code", "Station Name", "Ring", "IP", "Status"]

    def create_category_matrix_sheet(sheet_title: str, category_name: str, fill_style):
        ws = wb.create_sheet(title=sheet_title)
        param_list = sorted_cat_params[category_name]
        headers = base_cols + param_list

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = fill_style
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx, r in enumerate(results, start=2):
            s = r["site"]
            row_data = [
                s.get("no", ""),
                s.get("code", ""),
                s.get("name", ""),
                s.get("ring", ""),
                s.get("ip", ""),
                r["status"]
            ]
            c_dict = r["settings"].get(category_name, {}) if r["status"] == "ONLINE" else {}
            for p in param_list:
                row_data.append(c_dict.get(p, "-"))
            ws.append(row_data)

            # Style data cells
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=col_idx)
                c.border = thin_border
                if col_idx in [1, 2, 4, 6]:
                    c.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "G2"
        ws.auto_filter.ref = ws.dimensions

        # Adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col[:15])
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet: Battery Charge
    create_category_matrix_sheet("Battery_Charge", "Battery Charge", header_fill_brand)
    # Sheet: Rectifiers
    create_category_matrix_sheet("Rectifiers", "Rectifiers", header_fill_sub)
    # Sheet: Temperature
    create_category_matrix_sheet("Temperature", "Temperature", PatternFill(start_color="375623", end_color="375623", fill_type="solid"))
    # Sheet: System
    create_category_matrix_sheet("System", "System", PatternFill(start_color="595959", end_color="595959", fill_type="solid"))

    # ----------------------------------------------------
    # Sheet 1: All_Parameters_Raw (PIVOTED MATRIX: 1 ROW PER STATION)
    # Category + Parameter Name as Columns
    # ----------------------------------------------------
    ws_all_pivoted = wb.create_sheet(title="All_Parameters_Raw", index=0)

    # Collect all (Category, Parameter Name) in ordered sequence
    cat_fills = {
        "Battery Charge": header_fill_brand, # MEA Orange
        "Rectifiers": header_fill_sub,        # Dark Blue
        "Temperature": PatternFill(start_color="375623", end_color="375623", fill_type="solid"), # Forest Green
        "System": PatternFill(start_color="595959", end_color="595959", fill_type="solid")       # Gray
    }

    all_cat_param_pairs = []
    for c in cat_keys:
        for p in sorted_cat_params[c]:
            all_cat_param_pairs.append((c, p))

    all_pivoted_headers = base_cols + [f"[{c}] {p}" for c, p in all_cat_param_pairs]
    ws_all_pivoted.append(all_pivoted_headers)

    # Style Row 1 Headers with category color coding
    base_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    for col_idx in range(1, len(base_cols) + 1):
        cell = ws_all_pivoted.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = base_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, (c, p) in enumerate(all_cat_param_pairs, start=len(base_cols) + 1):
        cell = ws_all_pivoted.cell(row=1, column=idx)
        cell.font = header_font
        cell.fill = cat_fills.get(c, base_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Populate 1 row per station
    for r_idx, r in enumerate(results, start=2):
        s = r["site"]
        row_data = [
            s.get("no", ""),
            s.get("code", ""),
            s.get("name", ""),
            s.get("ring", ""),
            s.get("ip", ""),
            r["status"]
        ]
        settings_dict = r.get("settings", {}) if r["status"] == "ONLINE" else {}
        for c, p in all_cat_param_pairs:
            row_data.append(settings_dict.get(c, {}).get(p, "-"))
        ws_all_pivoted.append(row_data)

        # Style row data
        for col_idx in range(1, len(all_pivoted_headers) + 1):
            cell = ws_all_pivoted.cell(row=r_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 2, 4, 6]:
                cell.alignment = Alignment(horizontal="center")

    ws_all_pivoted.freeze_panes = "G2"
    ws_all_pivoted.auto_filter.ref = ws_all_pivoted.dimensions
    for col in ws_all_pivoted.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col[:15])
        ws_all_pivoted.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # ----------------------------------------------------
    # Sheet: All Parameters Long Format (Raw Key-Value Table)
    # ----------------------------------------------------
    ws_long = wb.create_sheet(title="All_Parameters_List")
    long_headers = [
        "No.", "Code", "Station Name", "Ring", "IP",
        "Category", "Parameter Name", "Value", "Unit / Type",
        "Raw Value", "Value Type", "Equip ID", "Signal Type", "Signal ID"
    ]
    ws_long.append(long_headers)
    for col_idx in range(1, len(long_headers) + 1):
        cell = ws_long.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill_brand
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        s = r["site"]
        if r["status"] == "ONLINE":
            for p in r["all_params"]:
                ws_long.append([
                    s.get("no", ""),
                    s.get("code", ""),
                    s.get("name", ""),
                    s.get("ring", ""),
                    s.get("ip", ""),
                    p.get("category", ""),
                    p.get("name", ""),
                    p.get("interpreted_value", ""),
                    p.get("unit", ""),
                    p.get("raw_value", ""),
                    p.get("val_type", ""),
                    p.get("equip_id", ""),
                    p.get("sig_type", ""),
                    p.get("sig_id", "")
                ])

    ws_long.freeze_panes = "A2"
    ws_long.auto_filter.ref = ws_long.dimensions
    for col in ws_long.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col[:15])
        ws_long.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 3. Save Excel Report
    ts_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    reports_dir = PROJECT_ROOT / "reports"
    archive_dir = reports_dir / "archive"
    reports_dir.mkdir(parents=True, exist_ok=True)
    archive_dir.mkdir(parents=True, exist_ok=True)

    out_latest = reports_dir / "Rectifier_Settings_All_Sites_Latest.xlsx"
    out_archive = archive_dir / f"Rectifier_Settings_All_Sites_{ts_str}.xlsx"

    # SMB drive safe saving via local temp
    temp_wb_path = Path(r"C:\Users\Prajuktra\AppData\Local\Temp") / f"settings_dump_{ts_str}.xlsx"
    wb.save(temp_wb_path)
    wb.close()

    import shutil
    shutil.copyfile(temp_wb_path, out_latest)
    shutil.copyfile(temp_wb_path, out_archive)

    print(f"[+] Excel reports saved successfully:")
    print(f"    - Latest  : {out_latest}")
    print(f"    - Archive : {out_archive}")
    print("=" * 70)


if __name__ == "__main__":
    main()
