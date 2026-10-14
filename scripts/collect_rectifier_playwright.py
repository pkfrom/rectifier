from __future__ import annotations

import argparse
import ast
import json
import re
import time
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, Browser, BrowserContext, Page, TimeoutError as PlaywrightTimeoutError


DATA_MARKER = '"%__data_end__%"'
SUMMARY_HEADERS = [
    "สถานะเก็บข้อมูล",
    "เวลาเก็บข้อมูล",
    "Site Name (หน้าเว็บ)",
    "System Name",
    # --- DC Output & Load ---
    "Output Voltage (V)",
    "Output Current (A)",
    "System Load (A)",
    "System Power (kW)",
    # --- AC Input ---
    "AC Phase A Voltage (V)",
    "AC Phase B Voltage (V)",
    "AC Phase C Voltage (V)",
    "AC Phase A Current (A)",
    "AC Phase B Current (A)",
    "AC Phase C Current (A)",
    # --- Battery ---
    "Battery Status",
    "Battery Current (A)",
    "Battery Capacity (%)",
    "Battery Temp (°C)",
    "Estimated Backup Time (h)",
    # --- Temperature & Environment ---
    "Ambient Temp (°C)",
    # --- Rectifier System ---
    "Number of Rectifiers",
    "Rectifier Total Current (A)",
    "Rectifier Used Capacity (%)",
    "Rectifier Max Capacity (%)",
    "Rectifier Avg Voltage (V)",
    # --- Controller Details ---
    "Controller Product Model",
    "Controller Serial Number",
    "Controller Hardware Version",
    "Controller Software Version",
    "Controller Config Version",
    "จำนวนอุปกรณ์ใน Inventory",
    "หน้าเว็บ",
    "รายละเอียดข้อผิดพลาด",
]
INVENTORY_HEADERS = [
    "No.",
    "ชื่อย่อ",
    "สถานีไฟฟ้า",
    "Ring",
    "IP",
    "Equipment",
    "Product Model",
    "Hardware Revision",
    "Serial Number",
    "Software Revision",
    "สถานะเก็บข้อมูล",
    "เวลาเก็บข้อมูล",
]
SETTINGS_HEADERS = [
    "No.",
    "ชื่อย่อ",
    "สถานีไฟฟ้า",
    "Ring",
    "IP",
    "สถานะเก็บข้อมูล",
    "เวลาเก็บข้อมูล",
    # --- Battery Charging Settings ---
    "Float Charge Voltage (V)",
    "Equalize Charge Voltage (V)",
    "Battery Current Limit (C10)",
    "Over Current Limit (C10)",
    "Batt Rated Capacity (Ah)",
    "Low Capacity Alarm (%)",
    "Auto/Manual State",
    "Automatic Equalize",
    "Cyclic Equalize",
    "Max Equalize Time (min)",
    "Temp Comp Center (°C)",
    "Temp Comp Slope (mV/°C)",
    "Over Voltage 1 (V)",
    "Over Voltage 2 (V)",
    "Under Voltage 1 (V)",
    "Under Voltage 2 (V)",
    # --- LVD Settings ---
    "LVD 1 Status",
    "LVD 1 Mode",
    "LVD 1 Disconnect Voltage (V)",
    "LVD 1 Reconnect Voltage (V)",
    "LVD 1 Reconnect Delay (min)",
    "LVD 2 Status",
    "LVD 2 Mode",
    "LVD 2 Disconnect Voltage (V)",
    "LVD 2 Reconnect Voltage (V)",
    "LVD 2 Reconnect Delay (min)",
    # --- Temperature Alarm Settings ---
    "Ambient Temp High Alarm 1 (°C)",
    "Ambient Temp High Alarm 2 (°C)",
    "Ambient Temp Low Alarm (°C)",
    "Battery Temp High Alarm 1 (°C)",
    "Battery Temp High Alarm 2 (°C)",
    "Temp Comp Probe Source",
    # --- Rectifier Settings ---
    "HVSD Limit (V)",
    "Walk-In Status",
    "Walk-In Time (s)",
    "Rectifier Current Limit (%)",
    "Default Voltage (V)",
    "ECO Mode",
    # --- Battery Test Settings ---
    "Batt Test Voltage Level (V)",
    "Batt End Test Voltage (V)",
    "Batt End Test Time (min)",
    "Batt End Test Capacity (%)",
]


def clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def parse_device_data(raw: bytes) -> dict[str, Any]:
    text = raw.decode("utf-8-sig", errors="replace")
    text = re.sub(r'/\*[\s\S]*?\*/|//.*', '', text).strip()
    text = text.split(DATA_MARKER, 1)[0].strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        python_text = re.sub(r"\btrue\b", "True", text)
        python_text = re.sub(r"\bfalse\b", "False", python_text)
        python_text = re.sub(r"\bnull\b", "None", python_text)
        python_text = re.sub(r"\bundefined\b", "None", python_text)
        python_text = re.sub(r"\bnan\b", "None", python_text, flags=re.IGNORECASE)
        parsed = ast.literal_eval(python_text)
        if not isinstance(parsed, dict):
            raise ValueError("ข้อมูลจากอุปกรณ์ไม่ได้เป็น object")
        return parsed


def extract_settings_dict(data: dict[str, Any]) -> dict[str, str]:
    if not isinstance(data, dict):
        return {}
    settings: dict[str, str] = {}
    for section_name, section_obj in data.items():
        if not isinstance(section_obj, dict):
            continue
        lst = section_obj.get("list", [])
        for item in lst:
            if isinstance(item, list) and len(item) >= 10:
                name = item[0][0] if isinstance(item[0], list) else str(item[0])
                name = name.strip()
                val_type = item[8]
                val = item[9]
                unit = item[10] if len(item) > 10 else ""
                
                if val_type == 5 and isinstance(unit, list):
                    try:
                        idx = int(val)
                        if 0 <= idx < len(unit):
                            val = unit[idx][0] if isinstance(unit[idx], list) else unit[idx]
                    except Exception:
                        pass
                elif isinstance(val, (int, float)):
                    val = f"{val:g}"
                elif isinstance(val, str):
                    val = val.strip()
                
                settings[f"{section_name}:{name}"] = str(val)
    return settings


def get_json(ip: str, path: str, timeout: float) -> dict[str, Any]:
    url = f"http://{ip}{path}?_={int(datetime.now().timestamp() * 1000)}"
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "MEA-Rectifier-Playwright/1.0", "Cache-Control": "no-cache"},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return parse_device_data(response.read())


def nested_first(value: Any, default: Any = "") -> Any:
    while isinstance(value, list) and value:
        value = value[0]
    return clean(value) if value is not None else default


def parse_telemetry_num(raw: Any) -> str:
    if isinstance(raw, list) and len(raw) >= 4:
        val, alarm, unvalid, not_config = raw[0], raw[1], raw[2], raw[3]
        if unvalid == 1 or not_config == 1:
            return ""
        if isinstance(val, list) and len(val) > 0:
            val = val[0]
        val_str = str(val).strip()
        m = re.search(r"[-+]?\d+(?:\.\d+)?", val_str)
        return f"{float(m.group(0)):g}" if m else val_str
    elif isinstance(raw, list) and len(raw) > 0:
        val = raw[0]
        if isinstance(val, list) and len(val) > 0:
            val = val[0]
        val_str = str(val).strip()
        m = re.search(r"[-+]?\d+(?:\.\d+)?", val_str)
        return f"{float(m.group(0)):g}" if m else val_str
    elif isinstance(raw, (int, float)):
        return f"{raw:g}"
    elif isinstance(raw, str):
        val_str = raw.strip()
        m = re.search(r"[-+]?\d+(?:\.\d+)?", val_str)
        return f"{float(m.group(0)):g}" if m else val_str
    return ""


def parse_telemetry_text(raw: Any) -> str:
    if isinstance(raw, list) and len(raw) >= 4:
        val, alarm, unvalid, not_config = raw[0], raw[1], raw[2], raw[3]
        if unvalid == 1 or not_config == 1:
            return ""
        if isinstance(val, list) and len(val) > 0:
            val = val[0]
        return str(val).strip()
    elif isinstance(raw, list) and len(raw) > 0:
        val = raw[0]
        if isinstance(val, list) and len(val) > 0:
            val = val[0]
        return str(val).strip()
    elif raw is not None:
        return str(raw).strip()
    return ""


def collect_site_playwright(
    context: BrowserContext,
    site: dict[str, Any],
    timeout: float = 15.0,
    page_delay: float = 1.0,
    screenshot_dir: Path | None = None,
) -> dict[str, Any]:
    ip = str(site["IP"]).strip()
    collected_at = datetime.now().astimezone().isoformat(timespec="seconds")
    result = {**site, "status": "ERROR", "collected_at": collected_at, "error": ""}
    errors: list[str] = []

    captured_data: dict[str, bytes] = {}

    page = context.new_page()
    page.set_default_timeout(int(timeout * 1000))

    def on_response(resp):
        try:
            url_str = resp.url
            if "/var/datas/" in url_str:
                for key in ["login", "system_inventory", "index", "status", "main", "system_battery", "system_rectifier", "system_udef", "setting"]:
                    if f"data.{key}.html" in url_str:
                        try:
                            body = resp.body()
                            if body:
                                captured_data[key] = body
                        except Exception:
                            pass
        except Exception:
            pass

    page.on("response", on_response)

    target_url = f"http://{ip}/app/www_user/html/eng/index.html"
    try:
        page.goto(target_url, wait_until="domcontentloaded", timeout=int(timeout * 1000))
        if page_delay > 0:
            time.sleep(page_delay)
    except PlaywrightTimeoutError:
        try:
            # Fallback to base IP url
            page.goto(f"http://{ip}/", wait_until="domcontentloaded", timeout=int(timeout * 1000))
            if page_delay > 0:
                time.sleep(page_delay)
        except Exception as exc:
            errors.append(f"page_load: {type(exc).__name__}: {exc}")
    except Exception as exc:
        errors.append(f"page_load: {type(exc).__name__}: {exc}")

    if screenshot_dir:
        try:
            screenshot_dir.mkdir(parents=True, exist_ok=True)
            shot_file = screenshot_dir / f"{site.get('No.', 'site')}_{ip.replace('.', '_')}.png"
            page.screenshot(path=str(shot_file), full_page=False)
        except Exception as exc:
            print(f"  [!] ไม่สามารถแคปเจอร์หน้าจอ: {exc}")

    try:
        page.close()
    except Exception:
        pass

    # Helper to retrieve data from captured network or direct HTTP fallback
    def fetch_data(label: str, key_or_paths: list[str] | str, required: bool = True) -> dict[str, Any]:
        keys = [key_or_paths] if isinstance(key_or_paths, str) else key_or_paths
        for k in keys:
            norm_key = k.replace("/var/datas/data.", "").replace(".html", "")
            if norm_key in captured_data:
                try:
                    return parse_device_data(captured_data[norm_key])
                except Exception:
                    pass

        # Fallback to direct HTTP fetch
        path_list = [key_or_paths] if isinstance(key_or_paths, str) else key_or_paths
        last_error: Exception | None = None
        for p in path_list:
            formatted_path = p if p.startswith("/") else f"/var/datas/data.{p}.html"
            for attempt in range(2):
                if attempt > 0:
                    time.sleep(0.3)
                try:
                    return get_json(ip, formatted_path, timeout=min(timeout, 5.0))
                except urllib.error.HTTPError as exc:
                    last_error = exc
                    if exc.code == 404:
                        break
                except Exception as exc:
                    last_error = exc
        if required and last_error is not None:
            errors.append(f"{label}: {type(last_error).__name__}: {last_error}")
        return {}

    login = fetch_data("login", "login", required=True)
    inventory_data = fetch_data("inventory", "system_inventory", required=True)
    index = fetch_data("status", ["index", "status", "main"], required=False)
    battery_data = fetch_data("battery", "system_battery", required=False)
    rectifier_data = fetch_data("rectifier", "system_rectifier", required=False)
    udef_data = fetch_data("udef", "system_udef", required=False)
    setting_data = fetch_data("setting", "setting", required=False)
    settings = extract_settings_dict(setting_data)

    inventory: list[dict[str, Any]] = []
    try:
        for item in inventory_data.get("inventory", []):
            if not isinstance(item, list) or len(item) < 5:
                continue
            equipment = item[0]
            if isinstance(equipment, list):
                equipment = equipment[0] if equipment else ""
            inventory.append(
                {
                    "equipment": clean(equipment),
                    "product_model": clean(item[1]),
                    "hardware_revision": clean(item[2]),
                    "serial_number": clean(item[3]),
                    "software_revision": clean(item[4]),
                }
            )

        main_content = index.get("mainContent", {})
        system = main_content.get("system", {})

        # --- DC Output & Load ---
        voltage = parse_telemetry_num(main_content.get("VAmeter", {}).get("value", {}).get("outer", ""))
        current = parse_telemetry_num(system.get("DC", ""))
        if not current:
            current = parse_telemetry_num(main_content.get("PAmeter", {}).get("value", {}).get("inner", ""))

        sys_load = current
        sys_power = ""
        udef_signal = udef_data.get("User_Defined", {}).get("SIGNAL", [])
        if len(udef_signal) >= 2 and isinstance(udef_signal[0], list) and isinstance(udef_signal[1], list):
            sig_names = [s[0] if isinstance(s, list) else str(s) for s in udef_signal[0]]
            for idx_s, name in enumerate(sig_names):
                if idx_s < len(udef_signal[1]):
                    s_val = udef_signal[1][idx_s]
                    if "load" in name.lower() and parse_telemetry_num(s_val):
                        sys_load = parse_telemetry_num(s_val)
                    elif "power" in name.lower() and "consumption" not in name.lower() and parse_telemetry_num(s_val):
                        sys_power = parse_telemetry_num(s_val)

        # --- AC Input ---
        mains = system.get("Mains", [])
        ac_curr = system.get("ACCurrent", [])
        ac_v_a = parse_telemetry_num(mains[0]) if len(mains) > 0 else ""
        ac_v_b = parse_telemetry_num(mains[1]) if len(mains) > 1 else ""
        ac_v_c = parse_telemetry_num(mains[2]) if len(mains) > 2 else ""
        ac_i_a = parse_telemetry_num(ac_curr[0]) if len(ac_curr) > 0 else ""
        ac_i_b = parse_telemetry_num(ac_curr[1]) if len(ac_curr) > 1 else ""
        ac_i_c = parse_telemetry_num(ac_curr[2]) if len(ac_curr) > 2 else ""

        # --- Battery ---
        batt_stat = parse_telemetry_text(system.get("Battery_status", ""))
        batt_curr = parse_telemetry_num(system.get("Battery", ""))
        batt_cap = parse_telemetry_num(system.get("BATT_CAPACITY", ""))
        batt_temp = parse_telemetry_num(battery_data.get("group", {}).get("chart", {}).get("Battery_Temperature", ""))
        batt_backup = parse_telemetry_num(battery_data.get("group", {}).get("chart", {}).get("Estimated_Remaining_Time", ""))

        # --- Temperature & Environment ---
        amb_temp = parse_telemetry_num(main_content.get("drawThermometer", {}).get("value", ""))

        # --- Rectifier System ---
        rect_curr = parse_telemetry_num(rectifier_data.get("group", {}).get("chart", {}).get("curr", ""))
        rect_cap = parse_telemetry_num(rectifier_data.get("group", {}).get("chart", {}).get("used_cap", ""))
        rect_list = rectifier_data.get("group", {}).get("list", [])
        rect_items = {}
        for item in rect_list:
            if isinstance(item, list) and len(item) >= 2:
                name_en = item[0][0] if isinstance(item[0], list) else str(item[0])
                rect_items[name_en] = parse_telemetry_num(item[1])

        rect_num = rect_items.get("Number of Rectifiers", "")
        if not rect_num:
            rect_num = sum(1 for row in inventory if str(row["equipment"]).lower().startswith("rectifier"))
        rect_avg_v = rect_items.get("Average Voltage", "")
        rect_max_cap = rect_items.get("Maximum Capacity Used", "")

        # --- Controller Details ---
        controller = index.get("sider", {}).get("system", [])
        if not controller:
            controller = [
                login.get("product_model", ""),
                login.get("serial_number", ""),
                login.get("hardware_version", ""),
                login.get("software_version", ""),
                login.get("config_version", ""),
            ]
        controller = list(controller) + [""] * (5 - len(controller))
        page_system_name = login.get("system_name", [""])
        page_site_name = login.get("site_name", [""])

        status = "OK" if inventory and not errors else ("PARTIAL" if inventory else "ERROR")
        result.update(
            {
                "status": status,
                "page_site_name": nested_first(page_site_name),
                "system_name": nested_first(page_system_name),
                # DC
                "output_voltage": voltage,
                "output_current": current,
                "system_load": sys_load,
                "system_power": sys_power,
                # AC
                "ac_v_a": ac_v_a,
                "ac_v_b": ac_v_b,
                "ac_v_c": ac_v_c,
                "ac_i_a": ac_i_a,
                "ac_i_b": ac_i_b,
                "ac_i_c": ac_i_c,
                # Battery
                "battery_status": batt_stat,
                "battery_current": batt_curr,
                "battery_capacity": batt_cap,
                "battery_temp": batt_temp,
                "battery_backup_hours": batt_backup,
                # Environment
                "ambient_temp": amb_temp,
                # Rectifiers
                "rectifier_count": rect_num,
                "rectifier_total_current": rect_curr,
                "rectifier_used_capacity": rect_cap,
                "rectifier_max_capacity": rect_max_cap,
                "rectifier_avg_voltage": rect_avg_v,
                # Controller
                "controller_product_model": clean(controller[0]),
                "controller_serial_number": clean(controller[1]),
                "controller_hardware_version": clean(controller[2]),
                "controller_software_version": clean(controller[3]),
                "controller_config_version": clean(controller[4]),
                "inventory": inventory,
                "settings": settings,
                "error": " | ".join(errors)[:500],
                "page_url": f"http://{ip}/app/www_user/html/eng/index.html",
            }
        )
    except (TypeError, IndexError, KeyError, ValueError) as exc:
        result["error"] = (" | ".join(errors + [f"processing: {type(exc).__name__}: {exc}"]))[:500]
        result["inventory"] = []
        result["page_url"] = f"http://{ip}/app/www_user/html/eng/index.html"
    return result


def load_sites(workbook_path: Path) -> tuple[Any, list[dict[str, Any]], int]:
    workbook = load_workbook(workbook_path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
    required = ["ชื่อย่อ", "สถานีไฟฟ้า", "Ring", "No.", "IP"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"ไม่พบคอลัมน์ที่ต้องใช้: {', '.join(missing)}")

    start_col = max(headers.values()) + 1
    sites = []
    for row in range(2, sheet.max_row + 1):
        ip = sheet.cell(row, headers["IP"]).value
        if ip is None or not str(ip).strip():
            continue
        sites.append(
            {
                "source_row": row,
                "ชื่อย่อ": sheet.cell(row, headers["ชื่อย่อ"]).value,
                "สถานีไฟฟ้า": sheet.cell(row, headers["สถานีไฟฟ้า"]).value,
                "Ring": sheet.cell(row, headers["Ring"]).value,
                "No.": sheet.cell(row, headers["No."]).value,
                "IP": str(ip).strip(),
            }
        )
    return workbook, sites, start_col


def style_header(sheet, row: int, start_column: int, end_column: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for col in range(start_column, end_column + 1):
        cell = sheet.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(sheet, min_width: int = 10, max_width: int = 45) -> None:
    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(length + 2, min_width), max_width
        )


def write_results(workbook, results: list[dict[str, Any]], output_path: Path, start_col: int = 6) -> None:
    summary = workbook[workbook.sheetnames[0]]
    for offset, header in enumerate(SUMMARY_HEADERS):
        summary.cell(1, start_col + offset, header)
    style_header(summary, 1, 1, start_col + len(SUMMARY_HEADERS) - 1)

    web_page_offset = SUMMARY_HEADERS.index("หน้าเว็บ")
    result_by_row = {item["source_row"]: item for item in results}
    for row in range(2, summary.max_row + 1):
        result = result_by_row.get(row)
        if not result:
            continue
        values = [
            result["status"],
            result["collected_at"],
            result.get("page_site_name", ""),
            result.get("system_name", ""),
            # DC
            result.get("output_voltage", ""),
            result.get("output_current", ""),
            result.get("system_load", ""),
            result.get("system_power", ""),
            # AC
            result.get("ac_v_a", ""),
            result.get("ac_v_b", ""),
            result.get("ac_v_c", ""),
            result.get("ac_i_a", ""),
            result.get("ac_i_b", ""),
            result.get("ac_i_c", ""),
            # Battery
            result.get("battery_status", ""),
            result.get("battery_current", ""),
            result.get("battery_capacity", ""),
            result.get("battery_temp", ""),
            result.get("battery_backup_hours", ""),
            # Environment
            result.get("ambient_temp", ""),
            # Rectifiers
            result.get("rectifier_count", ""),
            result.get("rectifier_total_current", ""),
            result.get("rectifier_used_capacity", ""),
            result.get("rectifier_max_capacity", ""),
            result.get("rectifier_avg_voltage", ""),
            # Controller
            result.get("controller_product_model", ""),
            result.get("controller_serial_number", ""),
            result.get("controller_hardware_version", ""),
            result.get("controller_software_version", ""),
            result.get("controller_config_version", ""),
            len(result.get("inventory", [])),
            result["page_url"],
            result.get("error", ""),
        ]
        for offset, value in enumerate(values):
            summary.cell(row, start_col + offset, value)
        link_cell = summary.cell(row, start_col + web_page_offset)
        link_cell.hyperlink = result["page_url"]
        link_cell.style = "Hyperlink"

    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.sheet_view.showGridLines = True
    autosize(summary)

    if "Settings" in workbook.sheetnames:
        del workbook["Settings"]
    settings_sheet = workbook.create_sheet("Settings")
    settings_sheet.append(SETTINGS_HEADERS)
    for result in results:
        s = result.get("settings", {})
        settings_sheet.append(
            [
                result["No."],
                result["ชื่อย่อ"],
                result["สถานีไฟฟ้า"],
                result["Ring"],
                result["IP"],
                result["status"],
                result["collected_at"],
                # --- Battery Charging Settings ---
                s.get("CHARGE:Float Charge Voltage", ""),
                s.get("CHARGE:Equalize Charge Voltage", ""),
                s.get("CHARGE:Battery Current Limit", ""),
                s.get("CHARGE:Over Current Limit", ""),
                s.get("CHARGE:Batt1 Rated Capacity", ""),
                s.get("CHARGE:Low Capacity Point", ""),
                s.get("CHARGE:Auto/Manual State", ""),
                s.get("CHARGE:Automatic Equalize", ""),
                s.get("CHARGE:Cyclic Equalize", ""),
                s.get("CHARGE:Maximum Equalize Charge Time", ""),
                s.get("CHARGE:Temperature Compensation Center", ""),
                s.get("CHARGE:Temp Comp Coefficient (slope)", ""),
                s.get("CHARGE:Over Voltage 1", ""),
                s.get("CHARGE:Over Voltage 2", ""),
                s.get("CHARGE:Under Voltage 1", ""),
                s.get("CHARGE:Under Voltage 2", ""),
                # --- LVD Settings ---
                s.get("LVD:LVD 1", ""),
                s.get("LVD:LVD 1 Mode", ""),
                s.get("LVD:LVD 1 Voltage", ""),
                s.get("LVD:LVD 1 Reconnect Voltage", ""),
                s.get("LVD:LVD 1 Reconnect Delay", ""),
                s.get("LVD:LVD 2", ""),
                s.get("LVD:LVD 2 Mode", ""),
                s.get("LVD:LVD 2 Voltage", ""),
                s.get("LVD:LVD 2 Reconnect Voltage", ""),
                s.get("LVD:LVD 2 Reconnect Delay", ""),
                # --- Temperature Alarm Settings ---
                s.get("TEMPERATURE:Ambient Temp High1", ""),
                s.get("TEMPERATURE:Ambient Temp High2", ""),
                s.get("TEMPERATURE:Ambient Temp Low", ""),
                s.get("TEMPERATURE:Comp Temp High1", ""),
                s.get("TEMPERATURE:Comp Temp High2", ""),
                s.get("TEMPERATURE:Temp Compensation Probe Number", ""),
                # --- Rectifier Settings ---
                s.get("RECT:HVSD Limit", ""),
                s.get("RECT:Walk-In", ""),
                s.get("RECT:Walk-In Time", ""),
                s.get("RECT:Rectifier Current Limit", ""),
                s.get("RECT:Default Voltage", ""),
                s.get("RECT:ECO Mode", s.get("ECO:ECO Mode", "")),
                # --- Battery Test Settings ---
                s.get("BATT_TEST:Test Voltage Level", ""),
                s.get("BATT_TEST:End Test Voltage", ""),
                s.get("BATT_TEST:End Test Time", ""),
                s.get("BATT_TEST:End Test Capacity", ""),
            ]
        )
    style_header(settings_sheet, 1, 1, len(SETTINGS_HEADERS))
    settings_sheet.freeze_panes = "A2"
    settings_sheet.auto_filter.ref = settings_sheet.dimensions
    settings_sheet.sheet_view.showGridLines = True
    autosize(settings_sheet)

    if "Inventory" in workbook.sheetnames:
        del workbook["Inventory"]
    inventory_sheet = workbook.create_sheet("Inventory")
    inventory_sheet.append(INVENTORY_HEADERS)
    for result in results:
        rows = result.get("inventory", [])
        if not rows:
            inventory_sheet.append(
                [
                    result["No."], result["ชื่อย่อ"], result["สถานีไฟฟ้า"], result["Ring"], result["IP"],
                    "", "", "", "", "", result["status"], result["collected_at"],
                ]
            )
            continue
        for item in rows:
            inventory_sheet.append(
                [
                    result["No."], result["ชื่อย่อ"], result["สถานีไฟฟ้า"], result["Ring"], result["IP"],
                    item["equipment"], item["product_model"], item["hardware_revision"],
                    item["serial_number"], item["software_revision"], result["status"], result["collected_at"],
                ]
            )
    style_header(inventory_sheet, 1, 1, len(INVENTORY_HEADERS))
    inventory_sheet.freeze_panes = "A2"
    inventory_sheet.auto_filter.ref = inventory_sheet.dimensions
    inventory_sheet.sheet_view.showGridLines = True
    autosize(inventory_sheet)

    workbook.save(output_path)


def main() -> int:
    base_dir = Path(__file__).resolve().parent.parent
    reports_dir = base_dir / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    default_in = (base_dir / "Rectifier.xlsx") if (base_dir / "Rectifier.xlsx").exists() else Path("Rectifier.xlsx")
    default_out = reports_dir / "Rectifier_Inventory_Playwright.xlsx"

    parser = argparse.ArgumentParser(description="เก็บข้อมูลหน้า Web Controller ของ Vertiv NetSure ผ่าน Playwright แบบเปิดเบราว์เซอร์")
    parser.add_argument("--input", type=Path, default=default_in, help="ไฟล์ Excel ต้นทาง")
    parser.add_argument("--output", type=Path, default=default_out, help="ไฟล์ Excel ปลายทาง")
    parser.add_argument("--timeout", type=float, default=12.0, help="Page/Request timeout (วินาที)")
    parser.add_argument("--ip", type=str, default=None, help="เก็บข้อมูลเฉพาะ IP ที่ระบุ (filter หรือทดสอบ)")
    parser.add_argument("--limit", type=int, default=None, help="จำกัดจำนวน site ที่ต้องการดึงข้อมูล")
    parser.add_argument("--headless", action="store_true", default=False, help="เปิดโหมด headless (ค่าเริ่มต้น: False เพื่อเปิดหน้าต่าง browser ให้ดู)")
    parser.add_argument("--slowmo", type=float, default=200.0, help="หน่วงเวลาการทำงานของ Browser แต่ละ Action (ms)")
    parser.add_argument("--delay", type=float, default=1.0, help="หน่วงเวลาหลังจากเปิดหน้าเว็บสำเร็จ (วินาที)")
    parser.add_argument("--screenshot", action="store_true", default=False, help="บันทึกภาพหน้าจอเก็บไว้ในโฟลเดอร์ screenshots/")
    args = parser.parse_args()

    workbook, sites, start_col = load_sites(args.input)

    if args.ip:
        target_ip = args.ip.strip()
        matched = [s for s in sites if s["IP"] == target_ip]
        if matched:
            sites = matched
        else:
            print(f"คำเตือน: ไม่พบ IP {target_ip} ใน {args.input}, สร้างรายการทดสอบชั่วคราว")
            sites = [{
                "source_row": 2,
                "ชื่อย่อ": "TEST",
                "สถานีไฟฟ้า": "TEST",
                "Ring": "TEST",
                "No.": 1,
                "IP": target_ip,
            }]

    if args.limit and args.limit > 0:
        sites = sites[:args.limit]

    screenshot_dir = Path("screenshots") if args.screenshot else None

    print(f"เริ่มดึงข้อมูลผ่าน Playwright ({len(sites)} sites)")
    print(f"Browser Mode: {'Headless' if args.headless else 'Headed (เปิดหน้าต่าง browser)'}")

    results: list[dict[str, Any]] = []
    with sync_playwright() as playwright:
        browser: Browser = playwright.chromium.launch(
            headless=args.headless,
            slow_mo=args.slowmo if not args.headless else 0,
            args=["--ignore-certificate-errors", "--disable-gpu", "--no-sandbox"]
        )
        context: BrowserContext = browser.new_context(
            viewport={"width": 1280, "height": 800},
            ignore_https_errors=True
        )

        try:
            for idx, site in enumerate(sites, start=1):
                ip = site["IP"]
                print(f"[{idx:03d}/{len(sites):03d}] กำลังเปิดหน้าเว็บ {ip} ...", end=" ", flush=True)
                result = collect_site_playwright(
                    context,
                    site,
                    timeout=args.timeout,
                    page_delay=args.delay,
                    screenshot_dir=screenshot_dir
                )
                results.append(result)
                print(f"[{result['status']}]")
        finally:
            context.close()
            browser.close()

    results.sort(key=lambda item: item["source_row"])
    write_results(workbook, results, args.output, start_col=start_col)
    ok_count = sum(item["status"] == "OK" for item in results)
    partial_count = sum(item["status"] == "PARTIAL" for item in results)
    error_count = sum(item["status"] == "ERROR" for item in results)
    print(
        f"เสร็จสิ้น: OK={ok_count}, PARTIAL={partial_count}, ERROR={error_count}, "
        f"output={args.output.resolve()}"
    )
    return 0 if error_count == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
