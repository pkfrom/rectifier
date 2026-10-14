from pathlib import Path
import openpyxl

base_dir = Path(__file__).resolve().parent.parent
reports_dir = base_dir / "reports"
wb_abnormal_path = (reports_dir / "Rectifier_Abnormal_Report_Latest.xlsx") if (reports_dir / "Rectifier_Abnormal_Report_Latest.xlsx").exists() else (reports_dir / "Rectifier_Abnormal_Report.xlsx")
wb = openpyxl.load_workbook(str(wb_abnormal_path))

# Update Sheet 1: 'การตั้งค่าผิดปกติ'
ws1 = wb['การตั้งค่าผิดปกติ']
for r in range(2, ws1.max_row + 1):
    code = str(ws1.cell(r, 1).value or "").strip()
    if code in ['AUG', 'KP', 'RTR', 'SYI']:
        ws1.cell(r, 8, value="53.6")
        note = str(ws1.cell(r, 10).value or "")
        note = note.replace(" | Equalize 56.4V (สูงเกิน)", " | Equalize ปรับเป็น 53.6V แล้ว")
        ws1.cell(r, 10, value=note)

# Update Sheet 2: 'การตั้งค่าทุกสถานี'
ws2 = wb['การตั้งค่าทุกสถานี']
eq_col = None
for c in range(1, ws2.max_column + 1):
    val = ws2.cell(1, c).value
    if val and 'Equalize Charge Voltage' in str(val):
        eq_col = c
        break

print("Sheet 2 Equalize col:", eq_col)
if eq_col:
    for r in range(2, ws2.max_row + 1):
        code = str(ws2.cell(r, 2).value or "").strip()
        if code in ['AUG', 'KP', 'RTR', 'SYI']:
            ws2.cell(r, eq_col, value="53.6")
            ws2.cell(r, eq_col).fill = openpyxl.styles.PatternFill(fill_type=None)
            print(f"Sheet 2: updated {code} row {r} -> 53.6")

# Update Sheet 0: 'สรุปภาพรวม'
ws0 = wb['สรุปภาพรวม']
for r in range(1, ws0.max_row + 1):
    for c in range(1, ws0.max_column + 1):
        cell_val = ws0.cell(r, c).value
        if cell_val and "56.4 V" in str(cell_val):
            new_val = str(cell_val).replace("56.4 V (AUG, KP, RTR, SYI)", "53.6 V (แก้ไขปรับตั้งค่า AUG, KP, RTR, SYI เรียบร้อยแล้ว)")
            ws0.cell(r, c, value=new_val)
            print(f"Updated summary at row {r}, col {c}")

wb.save(str(wb_abnormal_path))
print(f"Saved {wb_abnormal_path}")

# Update Rectifier_Inventory_188_Sites.xlsx
inv_path = reports_dir / 'Rectifier_Inventory_188_Sites.xlsx'
wb_inv = openpyxl.load_workbook(str(inv_path))
if "Settings" in wb_inv.sheetnames:
    ws_set = wb_inv["Settings"]
    eq_c = None
    for c in range(1, ws_set.max_column + 1):
        val = ws_set.cell(1, c).value
        if val and 'Equalize Charge Voltage' in str(val):
            eq_c = c
            break
    if eq_c:
        for r in range(2, ws_set.max_row + 1):
            code = str(ws_set.cell(r, 2).value or "").strip()
            if code in ['AUG', 'KP', 'RTR', 'SYI']:
                ws_set.cell(r, eq_c, value="53.6")
                ws_set.cell(r, eq_c).fill = openpyxl.styles.PatternFill(fill_type=None)
                print(f"Inventory: updated {code} row {r} -> 53.6")
    wb_inv.save(str(inv_path))
    print(f"Saved {inv_path}")
