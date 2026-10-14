import tarfile
import struct
import urllib.request
import json
import base64
import sys
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

from collect_rectifier_inventory import parse_device_data

# 1. Parse SettingParam.run
with tarfile.open('config/SettingParam.tar') as t:
    param_run_bytes = t.extractfile('./config/run/SettingParam.run').read()
    main_cfg_text = t.extractfile('./config/run/MainConfig.run').read().decode('utf-8', errors='replace')

offset = 24
records = []
while offset + 16 <= len(param_run_bytes):
    equip, sig_id, sig_type, val_int, ts = struct.unpack_from('<IHHiI', param_run_bytes, offset)
    val_float, = struct.unpack_from('<f', param_run_bytes, offset + 8)
    if equip > 1000 or sig_type not in (1, 2, 3, 4, 5):
        break
    records.append({
        'offset': offset,
        'equip': equip,
        'sig_type': sig_type,
        'sig_id': sig_id,
        'val_int': val_int,
        'val_float': val_float,
        'ts': ts
    })
    offset += 16

# Extract credentials at end of SettingParam.run
cred_str = param_run_bytes[1080:].decode('latin1', errors='ignore')
creds = re.findall(r'[a-zA-Z0-9]{3,}', cred_str)

# 2. Parse MainConfig.run
main_cfg_rows = []
for line in main_cfg_text.splitlines():
    line = line.strip()
    if not line or line.startswith('#') or line.startswith('['):
        continue
    parts = line.split('\t')
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) >= 7:
        main_cfg_rows.append(parts)

# 3. Fetch live endpoints from SE (10.209.57.21) and BAT (10.209.1.21)
endpoints = [
    '/var/datas/data.setting.html',
    '/var/datas/data.sys_setting.html',
    '/var/datas/data.udef_setting.html',
    '/var/datas/data.install_wizard.html',
    '/var/datas/data.system_rectifier.html',
    '/var/datas/data.system_battery.html',
    '/var/datas/data.system_dc.html',
    '/var/datas/data.system_ac.html',
    '/var/datas/data.hybrid.html',
    '/var/datas/data.alarms.html',
]

def fetch_all_signals(ip):
    signals = {}
    all_param_list = []
    
    def walk(ep, prefix, node):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(ep, f"{prefix}.{k}" if prefix else k, v)
        elif isinstance(node, list):
            for item in node:
                if isinstance(item, list) and len(item) >= 10 and isinstance(item[1], int) and isinstance(item[2], int) and isinstance(item[3], int):
                    name = item[0][0] if isinstance(item[0], list) else str(item[0])
                    eq, st, si = item[1], item[2], item[3]
                    vt = item[8]
                    cval = item[9]
                    opts = item[10] if len(item) > 10 else None
                    entry = {
                        'endpoint': ep,
                        'path': prefix,
                        'name': name,
                        'equip_id': eq,
                        'sig_type': st,
                        'sig_id': si,
                        'value_type': vt,
                        'live_val': cval,
                        'options': opts,
                        'raw': item
                    }
                    signals[(eq, st, si)] = entry
                    all_param_list.append(entry)
                elif isinstance(item, dict):
                    walk(ep, prefix, item)

    for ep in endpoints:
        url = f'http://{ip}{ep}'
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "MEA-Rectifier/1.0"})
            data = parse_device_data(urllib.request.urlopen(req, timeout=5).read())
            walk(ep, '', data)
        except Exception as e:
            pass
    return signals, all_param_list

print("Fetching live data from SE (10.209.57.21)...")
se_signals, se_all_params = fetch_all_signals('10.209.57.21')
print("Fetching live data from BAT (10.209.1.21)...")
bat_signals, bat_all_params = fetch_all_signals('10.209.1.21')

# Build Workbook
wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# ---------------- Sheet 1: SettingParam.run Analysis ----------------
ws1 = wb.create_sheet(title="SettingParam.run (Tar)")
headers1 = [
    "No.", "Equip ID", "Sig Type", "Sig ID", "Signal Key", 
    "Parameter Name", "Category / Section", "Value Type", 
    "Tar Value (Raw)", "Tar Value (Interpreted)", 
    "Live Value SE (10.209.57.21)", "Live Value BAT (10.209.1.21)", 
    "Tar Timestamp (Epoch)", "Tar Timestamp (Datetime)", "Options / Units"
]
ws1.append(headers1)

for i, r in enumerate(records, 1):
    key = (r['equip'], r['sig_type'], r['sig_id'])
    info = se_signals.get(key) or bat_signals.get(key)
    
    name = info['name'] if info else f"Unknown Signal [{r['equip']}:{r['sig_id']}]"
    path = info['path'] if info else "-"
    vt = info['value_type'] if info else (2 if abs(r['val_float']) > 0.001 and abs(r['val_float']) < 100000 else 3)
    opts = info['options'] if info else ""
    
    val_disp = round(r['val_float'], 2) if vt == 2 else r['val_int']
    opt_str = ""
    if vt == 5 and isinstance(opts, list):
        try:
            idx = int(r['val_int'])
            if 0 <= idx < len(opts):
                o = opts[idx]
                opt_str = f" ({o[0] if isinstance(o, list) else o})"
        except:
            pass
    elif vt == 2 and isinstance(opts, str):
        opt_str = f" {opts.strip()}"
        
    tar_val_interp = f"{val_disp}{opt_str}"
    
    se_val = se_signals.get(key, {}).get('live_val', '-')
    bat_val = bat_signals.get(key, {}).get('live_val', '-')
    
    # decode special known values if unmatched
    if not info:
        if r['equip'] == 1 and r['sig_id'] == 448:
            name = "Timezone Offset (Seconds)"
            path = "System Time"
            tar_val_interp = f"{r['val_int']} s (UTC+7)"
        elif r['equip'] == 1 and r['sig_id'] == 28:
            name = "NTP Sync Interval"
            path = "System Time"
            tar_val_interp = f"{r['val_int']} min"
        elif r['equip'] == 3 and r['sig_id'] == 1:
            name = "Rectifier #1 Phase"
            path = "Rectifier Phase"
            tar_val_interp = "Phase A (1)"
        elif r['equip'] == 4 and r['sig_id'] == 1:
            name = "Rectifier #2 Phase"
            path = "Rectifier Phase"
            tar_val_interp = "Phase B (2)"
        elif r['equip'] == 2 and r['sig_id'] == 18:
            name = "Rectifier Number in Group"
            path = "Rectifier System"
            tar_val_interp = f"{r['val_int']} units"

    try:
        dt_str = datetime.fromtimestamp(r['ts']).strftime('%Y-%m-%d %H:%M:%S')
    except:
        dt_str = str(r['ts'])

    row = [
        i, r['equip'], r['sig_type'], r['sig_id'], f"{r['equip']}:{r['sig_type']}:{r['sig_id']}",
        name, path, vt, r['val_int'], tar_val_interp,
        se_val, bat_val,
        r['ts'], dt_str, str(opts) if opts else ""
    ]
    ws1.append(row)

# ---------------- Sheet 2: MainConfig.run Analysis ----------------
ws2 = wb.create_sheet(title="MainConfig.run (Alarms & Site)")
headers2 = [
    "No.", "Timestamp", "Changed Info Type", "Text Type", "New Value / Text", 
    "Equip ID / Type ID", "Sig Type", "Sig ID", "Meaning / Description"
]
ws2.append(headers2)

for i, parts in enumerate(main_cfg_rows, 1):
    ts = parts[0]
    ctype = parts[1]
    ttype = parts[2]
    new_text = parts[3]
    eq = parts[4]
    st = parts[5]
    si = parts[6]
    
    desc = ""
    if ctype == "ST":
        desc = f"Site Name change to: {new_text}"
    elif ctype == "SL":
        desc = f"Site Location change to: {new_text}"
    elif ctype == "SD":
        desc = f"Site Description change to: {new_text}"
    elif ctype == "AL":
        desc = f"Alarm Level for Equip {eq} Signal {si} -> {new_text}"
    elif ctype == "AR":
        desc = f"Alarm Relay output mapping for Equip {eq} Signal {si} -> Relay {new_text}"
    else:
        desc = f"{ctype} change"
        
    ws2.append([i, ts, ctype, ttype, new_text, eq, st, si, desc])

# ---------------- Sheet 3: All Live Extractable Parameters ----------------
ws3 = wb.create_sheet(title="All Live Extractable Parameters")
headers3 = [
    "No.", "Endpoint", "Section / Path", "Parameter Name", 
    "Equip ID", "Sig Type", "Sig ID", "Signal Key", 
    "Value Type", "Value (SE - 10.209.57.21)", "Value (BAT - 10.209.1.21)", 
    "Unit / Options", "Settable via API"
]
ws3.append(headers3)

# deduplicate across live params
seen_keys = set()
count = 1
for p in se_all_params + bat_all_params:
    key = (p['equip_id'], p['sig_type'], p['sig_id'], p['name'])
    if key in seen_keys:
        continue
    seen_keys.add(key)
    
    sig_k = (p['equip_id'], p['sig_type'], p['sig_id'])
    se_val = se_signals.get(sig_k, {}).get('live_val', '-')
    bat_val = bat_signals.get(sig_k, {}).get('live_val', '-')
    
    settable = "Yes (POST web_cgi_control.cgi)" if p['sig_type'] in (1, 2) else "No (Telemetry/Status)"
    
    ws3.append([
        count, p['endpoint'], p['path'], p['name'],
        p['equip_id'], p['sig_type'], p['sig_id'], f"{p['equip_id']}:{p['sig_type']}:{p['sig_id']}",
        p['value_type'], se_val, bat_val,
        str(p['options']) if p['options'] else "", settable
    ])
    count += 1

# Format all sheets
for ws in [ws1, ws2, ws3]:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

out_file = "SettingParam_Extracted_Data.xlsx"
wb.save(out_file)
print(f"Saved report to {out_file}")
