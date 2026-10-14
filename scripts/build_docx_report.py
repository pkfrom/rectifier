import docx
import shutil
from pathlib import Path

base_dir = Path(__file__).resolve().parent.parent
reports_dir = base_dir / "reports"
main_doc = reports_dir / "Rectifier_Inspection_Report_All_Settings.docx"
doc = docx.Document(str(main_doc))

# Update paragraphs
for p in doc.paragraphs:
    if '56.4 V' in p.text or '56.4V' in p.text:
        p.text = p.text.replace('56.4 V (AUG, KP, RTR, SYI)', '53.6 V (ดำเนินการปรับตั้งค่า AUG, KP, RTR, SYI ให้เท่ากับ 53.6 V เรียบร้อยแล้ว)')
        p.text = p.text.replace('ลด Equalize Voltage ใน 4 สถานี (AUG, KP, RTR, SYI) จาก 56.4 V ลงมาเป็น 53.6 V', 'ดำเนินการปรับลด Equalize Voltage ใน 4 สถานี (AUG, KP, RTR, SYI) จาก 56.4 V เป็น 53.6 V เรียบร้อยแล้วทุกสถานี')

# Update tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '56.4' in cell.text:
                cell.text = cell.text.replace('56.4 V', '53.6 V (ปรับแล้ว)')
                cell.text = cell.text.replace('56.4', '53.6')
                cell.text = cell.text.replace('Equalize 56.4V (สูงเกิน)', 'Equalize ปรับเป็น 53.6V แล้ว')

# Save updated docx
out_path = reports_dir / 'Rectifier_Inspection_Report_All_Settings_Updated.docx'
doc.save(str(out_path))
print(f"Saved {out_path} successfully")

# Try to overwrite main docx if not locked
try:
    shutil.copy2(str(out_path), str(main_doc))
    print(f"Overwrote {main_doc} successfully")
except Exception as e:
    print("Main docx locked:", e)
