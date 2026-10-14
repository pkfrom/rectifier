import os
import sys
import re
import ast
import json
import time
import base64
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

DATA_MARKER = '"%__data_end__%"'

def parse_device_data(raw: bytes) -> dict:
    text = raw.decode("utf-8-sig", errors="replace")
    text = re.sub(r'/\*[\s\S]*?\*/|//.*', '', text).strip()
    text = text.split(DATA_MARKER, 1)[0].strip()
    try:
        return json.loads(text)
    except Exception:
        python_text = re.sub(r"\btrue\b", "True", text)
        python_text = re.sub(r"\bfalse\b", "False", python_text)
        python_text = re.sub(r"\bnull\b", "None", python_text)
        python_text = re.sub(r"\bundefined\b", "None", python_text)
        python_text = re.sub(r"\bnan\b", "None", python_text, flags=re.IGNORECASE)
        try:
            parsed = ast.literal_eval(python_text)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            pass
    return {}

ENDPOINTS = [
    ('setting', '/var/datas/data.setting.html'),
    ('sys_setting', '/var/datas/data.sys_setting.html'),
    ('udef_setting', '/var/datas/data.udef_setting.html'),
    ('install_wizard', '/var/datas/data.install_wizard.html'),
    ('system_inventory', '/var/datas/data.system_inventory.html'),
    ('system_rectifier', '/var/datas/data.system_rectifier.html'),
    ('system_battery', '/var/datas/data.system_battery.html'),
    ('system_dc', '/var/datas/data.system_dc.html'),
    ('system_ac', '/var/datas/data.system_ac.html'),
    ('hybrid', '/var/datas/data.hybrid.html'),
]

def fetch_endpoint(ip: str, path: str, timeout: float = 4.0) -> dict:
    url = f"http://{ip}{path}?_={int(time.time() * 1000)}"
    req = urllib.request.Request(url, headers={"User-Agent": "MEA-Rectifier/1.0", "Cache-Control": "no-cache"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return parse_device_data(resp.read())
    except Exception:
        return {}

def extract_parameters_from_dict(ep_name: str, data: dict) -> list:
    params = []
    def walk(prefix, node):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(f"{prefix}.{k}" if prefix else k, v)
        elif isinstance(node, list):
            for item in node:
                if isinstance(item, list) and len(item) >= 10 and isinstance(item[1], int) and isinstance(item[2], int) and isinstance(item[3], int):
                    name = item[0][0] if isinstance(item[0], list) else str(item[0])
                    eq, st, si = item[1], item[2], item[3]
                    vt = item[8]
                    cval = item[9]
                    opts = item[10] if len(item) > 10 else None
                    
                    opt_str = ""
                    if vt == 5 and isinstance(opts, list):
                        try:
                            idx = int(cval)
                            if 0 <= idx < len(opts):
                                o = opts[idx]
                                opt_str = o[0] if isinstance(o, list) else str(o)
                        except:
                            pass
                    elif vt == 2 and isinstance(opts, str):
                        opt_str = opts.strip()

                    params.append({
                        'endpoint': ep_name,
                        'path': prefix,
                        'name': name.strip(),
                        'equip_id': eq,
                        'sig_type': st,
                        'sig_id': si,
                        'value_type': vt,
                        'raw_value': cval,
                        'interpreted_value': opt_str if opt_str else str(cval),
                        'options_or_unit': str(opts) if opts else ""
                    })
                elif isinstance(item, dict):
                    walk(prefix, item)

    walk('', data)
    return params

def collect_site_all_params(site: dict) -> dict:
    code = site.get('code', '')
    name = site.get('name', '')
    ip = str(site.get('ip', '')).strip()
    ring = site.get('ring', '')
    no = site.get('no', '')

    if not ip or ip.lower() in ('none', 'nan', 'ไม่มี', '-'):
        return {
            'site': site,
            'status': 'NO_IP',
            'params': [],
            'inventory': [],
            'error': 'No IP Address'
        }

    site_params = []
    inv_list = []
    hw_info = {}

    for ep_name, ep_path in ENDPOINTS:
        d = fetch_endpoint(ip, ep_path)
        if not d:
            continue
        if ep_name == 'system_inventory':
            inv_list = d.get('inventory', [])
        
        extracted = extract_parameters_from_dict(ep_name, d)
        site_params.extend(extracted)

    status = 'OK' if site_params else 'ERROR'
    return {
        'site': site,
        'status': status,
        'params': site_params,
        'inventory': inv_list,
        'error': '' if status == 'OK' else 'Unable to fetch parameters'
    }

def main():
    base_dir = Path(__file__).resolve().parent.parent
    reports_dir = base_dir / "reports"
    archive_dir = reports_dir / "archive"
    reports_dir.mkdir(parents=True, exist_ok=True)
    archive_dir.mkdir(parents=True, exist_ok=True)

    input_file = (base_dir / 'Rectifier.xlsx') if (base_dir / 'Rectifier.xlsx').exists() else Path('Rectifier.xlsx')
    print(f"Reading site list from {input_file}...")
    wb_in = openpyxl.load_workbook(str(input_file), read_only=True)
    ws_in = wb_in.active

    sites = []
    for row in ws_in.iter_rows(values_only=True):
        if not row or not row[0] or str(row[0]).strip().lower() in ('ชื่อย่อ', 'code', 'station'):
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
        ring = row[2] if len(row) > 2 else ''
        no = row[3] if len(row) > 3 else ''
        ip = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ''
        sites.append({'code': code, 'name': name, 'ring': ring, 'no': no, 'ip': ip})

    print(f"Loaded {len(sites)} sites. Starting concurrent collection with 16 workers...")
    start_time = time.time()

    results = []
    with ThreadPoolExecutor(max_workers=16) as executor:
        futures = {executor.submit(collect_site_all_params, s): s for s in sites}
        done_count = 0
        for f in as_completed(futures):
            done_count += 1
            res = f.result()
            results.append(res)
            s = res['site']
            print(f"[{done_count:03d}/{len(sites):03d}] {s['code']:<5} {s['ip']:<15} {res['status']:<6} ({len(res['params'])} params)")

    elapsed = time.time() - start_time
    ok_count = sum(1 for r in results if r['status'] == 'OK')
    err_count = sum(1 for r in results if r['status'] != 'OK')
    print(f"\nCollection finished in {elapsed:.1f}s. OK={ok_count}, Failed/No_IP={err_count}")

    # Sort results by original site order
    site_order = {s['code']: i for i, s in enumerate(sites)}
    results.sort(key=lambda r: site_order.get(r['site']['code'], 999))

    # Build Excel Workbook
    timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_filename = str(archive_dir / f"Rectifier_All_Parameters_All_Sites_{timestamp_str}.xlsx")
    latest_filename = str(reports_dir / "Rectifier_All_Parameters_All_Sites_Latest.xlsx")

    print(f"Building Excel workbook {output_filename}...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    # 1. Summary Sheet
    ws_sum = wb.create_sheet(title="Station Summary")
    ws_sum.append(["No.", "Station Code", "Station Name", "Ring", "IP", "Status", "Param Count", "Inventory Items", "Error Note"])
    for idx, r in enumerate(results, 1):
        s = r['site']
        ws_sum.append([
            idx, s['code'], s['name'], s['ring'], s['ip'],
            r['status'], len(r['params']), len(r['inventory']), r['error']
        ])

    # 2. Main Settings Matrix Sheet (Key comparison table across all stations)
    # Collect all unique (path, name) across all sites
    param_keys = []
    seen_param_keys = set()
    for r in results:
        for p in r['params']:
            key = f"{p['path']}:{p['name']}"
            if key not in seen_param_keys:
                seen_param_keys.add(key)
                param_keys.append((key, p['name'], p['path'], p['equip_id'], p['sig_type'], p['sig_id']))

    # Sort param keys by path then name
    param_keys.sort(key=lambda x: (x[2], x[1]))

    ws_mat = wb.create_sheet(title="Settings Matrix (All Sites)")
    matrix_headers = ["Code", "Name", "IP"] + [f"{k[1]}" for k in param_keys]
    ws_mat.append(matrix_headers)

    for r in results:
        s = r['site']
        row = [s['code'], s['name'], s['ip']]
        site_pdict = {f"{p['path']}:{p['name']}": p['interpreted_value'] for p in r['params']}
        for k in param_keys:
            row.append(site_pdict.get(k[0], "-"))
        ws_mat.append(row)

    # 3. All Parameters Long Format (Every single parameter row)
    ws_all = wb.create_sheet(title="All Parameters (Raw)")
    ws_all.append([
        "Station Code", "Station Name", "IP", "Endpoint", "Path / Section", 
        "Parameter Name", "Equip ID", "Sig Type", "Sig ID", "Signal Key", 
        "Value Type", "Raw Value", "Interpreted Value", "Options / Unit"
    ])

    for r in results:
        s = r['site']
        for p in r['params']:
            ws_all.append([
                s['code'], s['name'], s['ip'], p['endpoint'], p['path'],
                p['name'], p['equip_id'], p['sig_type'], p['sig_id'],
                f"{p['equip_id']}:{p['sig_type']}:{p['sig_id']}",
                p['value_type'], p['raw_value'], p['interpreted_value'], p['options_or_unit']
            ])

    # 4. Inventory Sheet
    ws_inv = wb.create_sheet(title="Hardware Inventory")
    ws_inv.append(["Station Code", "Station Name", "IP", "Device Name", "Model", "Hardware Ver", "Serial Number", "Software Ver"])
    for r in results:
        s = r['site']
        for item in r['inventory']:
            if isinstance(item, list) and len(item) >= 5:
                d_name = item[0][0] if isinstance(item[0], list) else str(item[0])
                model = str(item[1]).strip()
                hw_ver = str(item[2]).strip()
                sn = str(item[3]).strip()
                sw_ver = str(item[4]).strip()
                ws_inv.append([s['code'], s['name'], s['ip'], d_name.strip(), model, hw_ver, sn, sw_ver])

    # Style header row for all sheets
    for ws in [ws_sum, ws_mat, ws_all, ws_inv]:
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    print("Saving Excel file...")
    wb.save(output_filename)
    wb.save(latest_filename)
    print(f"Successfully saved:")
    print(f"  - {output_filename}")
    print(f"  - {latest_filename}")

if __name__ == '__main__':
    main()
